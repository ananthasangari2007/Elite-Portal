from datetime import datetime, timedelta
from io import BytesIO

from flask import Blueprint, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user
from flask_wtf import FlaskForm
from sqlalchemy import func
from wtforms import DateField, SelectField, SubmitField, TimeField
from wtforms.validators import DataRequired

from app import db
from app.models import EliteSprintBid, EliteSprintSession, Task, User, PointTransaction
from app.services.elite_sprint import (
    APP_TIMEZONE,
    active_task_ids_by_type,
    get_active_session,
    latest_session as get_latest_session,
    local_datetime_to_utc,
    sprint_leaderboard,
    sprint_metrics,
    utc_datetime_to_local,
)
from app.utils.auth import role_required


elite_sprint_bp = Blueprint("elite_sprint", __name__, url_prefix="/elite-sprint-bidding")


class SprintSessionForm(FlaskForm):
    sprint_type = SelectField(
        "Sprint type",
        choices=[
            ("overall", "Overall Sprint"),
            ("daily", "Daily Sprint"),
            ("weekly", "Weekly Sprint"),
            ("monthly", "Monthly Sprint"),
        ],
        validators=[DataRequired()],
    )
    sprint_date = DateField("Sprint date", validators=[DataRequired()])
    start_time = TimeField("Start time", validators=[DataRequired()])
    end_time = TimeField("End time", validators=[DataRequired()])
    submit = SubmitField("Open Sprint")


@elite_sprint_bp.route("/")
@role_required("admin", "student")
def index():
    if current_user.role == "admin":
        return redirect(url_for("elite_sprint.admin"))
    return redirect(url_for("elite_sprint.student"))


@elite_sprint_bp.route("/admin", methods=["GET", "POST"])
@role_required("admin")
def admin():
    form = SprintSessionForm()
    if form.validate_on_submit():
        local_start_at = datetime.combine(form.sprint_date.data, form.start_time.data)
        local_end_at = datetime.combine(form.sprint_date.data, form.end_time.data)
        if local_end_at <= local_start_at:
            local_end_at = local_end_at + timedelta(days=1)
        start_at = local_datetime_to_utc(local_start_at)
        end_at = local_datetime_to_utc(local_end_at)
        if end_at <= datetime.utcnow():
            flash("End time must be in the future.", "danger")
        else:
            EliteSprintSession.query.filter_by(status="active").update({"status": "closed"})
            session = EliteSprintSession(
                sprint_date=form.sprint_date.data,
                sprint_type=form.sprint_type.data,
                start_time=start_at,
                end_time=end_at,
                status="active",
                created_by=current_user.id,
            )
            db.session.add(session)
            db.session.commit()
            flash("Elite Sprint Bidding is open for the selected window.", "success")
            return redirect(url_for("elite_sprint.admin"))

    active_session = get_active_session()
    latest_session = active_session or get_latest_session()
    return render_template(
        "elite_sprint/admin.html",
        form=form,
        active_session=active_session,
        latest_session=latest_session,
        timezone_label=APP_TIMEZONE,
        latest_start_local=utc_datetime_to_local(latest_session.start_time) if latest_session else None,
        latest_end_local=utc_datetime_to_local(latest_session.end_time) if latest_session else None,
        metrics=sprint_metrics(latest_session),
        leaderboard=sprint_leaderboard(latest_session.id if latest_session else None),
    )


@elite_sprint_bp.route("/admin/close/<int:session_id>", methods=["POST"])
@role_required("admin")
def close(session_id):
    session = EliteSprintSession.query.get_or_404(session_id)
    session.status = "closed"
    db.session.commit()
    leaders = sprint_leaderboard(session.id, limit=1)
    if leaders:
        flash(f"Elite Sprint Bidding has been closed. Highest bidder: {leaders[0]['student_name']} with {leaders[0]['total_tasks']} tasks.", "info")
    else:
        flash("Elite Sprint Bidding has been closed. No bids were submitted.", "info")
    return redirect(url_for("elite_sprint.admin"))


@elite_sprint_bp.route("/admin/verify-sprint/<int:session_id>", methods=["POST"])
@role_required("admin")
def verify_sprint(session_id):
    from app.services.points import record_penalty
    session = EliteSprintSession.query.get_or_404(session_id)
    bids = EliteSprintBid.query.filter_by(session_id=session.id).all()

    verified_count = 0
    penalty_count = 0
    star_count = 0

    for bid in bids:
        student = bid.student

        # Gather all bidded task IDs (lowercase or uppercase, stripped)
        daily_bids = [t.strip().upper() for t in (bid.daily_tasks or []) if t.strip()]
        weekly_bids = [t.strip().upper() for t in (bid.weekly_tasks or []) if t.strip()]
        monthly_bids = [t.strip().upper() for t in (bid.monthly_tasks or []) if t.strip()]
        all_bids = set(daily_bids + weekly_bids + monthly_bids)

        if not all_bids:
            bid.golden_star = False
            continue

        # Fetch approved task submissions for this student on the sprint date
        approved_subs = Submission.query.filter_by(
            student_id=student.id,
            submission_date=session.sprint_date,
            status="approved",
        ).all()

        # Build set of completed task codes/IDs
        completed_task_identifiers = set()
        for sub in approved_subs:
            completed_task_identifiers.add(str(sub.task_id))
            if sub.task.task_code:
                completed_task_identifiers.add(sub.task.task_code.strip().upper())

        # Find missed tasks
        missed_tasks = all_bids - completed_task_identifiers

        # Delete existing penalty transactions for this student on this session date
        # to ensure idempotency (running verification multiple times is safe)
        existing_penalties = PointTransaction.query.filter(
            PointTransaction.student_id == student.id,
            PointTransaction.type == "penalty",
            PointTransaction.reason.like(f"Elite Sprint Penalty (Session {session.id}):%"),
        ).all()
        for ep in existing_penalties:
            db.session.delete(ep)

        if not missed_tasks:
            bid.golden_star = True
            star_count += 1
        else:
            bid.golden_star = False
            # Deduct points of the missed bidded task IDs
            for missed_id in missed_tasks:
                # Find task
                task = Task.query.filter(
                    (Task.id == missed_id) | (func.upper(Task.task_code) == missed_id.upper())
                ).first()
                if task:
                    record_penalty(
                        student_id=student.id,
                        points=task.reward_points,
                        reason=f"Elite Sprint Penalty (Session {session.id}): Missed bidded task {task.task_code or task.id}",
                        approved_by=current_user.id,
                    )
            penalty_count += 1

        verified_count += 1

    # Flush deletes before querying active penalties
    db.session.flush()

    # Recalculate has_active_sprint_penalty for all processed users
    for bid in bids:
        student = bid.student
        any_sprint_penalties = PointTransaction.query.filter(
            PointTransaction.student_id == student.id,
            PointTransaction.type == "penalty",
            PointTransaction.reason.like("Elite Sprint Penalty%"),
        ).first()
        student.has_active_sprint_penalty = any_sprint_penalties is not None

    db.session.commit()
    flash(
        f"Verification complete for {verified_count} students. Golden Stars: {star_count}, Penalty Cases: {penalty_count}.",
        "success",
    )
    return redirect(url_for("elite_sprint.admin"))


@elite_sprint_bp.route("/admin/export-daily-report/<int:session_id>")
@role_required("admin")
def export_daily_report(session_id):
    import os
    from flask import current_app, send_file
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, PieChart, Reference

    session = EliteSprintSession.query.get_or_404(session_id)
    
    # 1. Determine folder and filename
    reports_dir = os.path.join(current_app.root_path, "..", "reports")
    os.makedirs(reports_dir, exist_ok=True)

    sprint_date = session.sprint_date
    month_name = sprint_date.strftime("%B").upper()
    year = sprint_date.year
    filename = f"{month_name}-{year} TASK TRACK SHEET.xlsx"
    filepath = os.path.join(reports_dir, filename)

    # 2. Load or create workbook
    if os.path.exists(filepath):
        workbook = load_workbook(filepath)
    else:
        workbook = Workbook()
        # Remove default sheet if it exists
        if "Sheet" in workbook.sheetnames:
            del workbook["Sheet"]

    # 3. Sheet names
    date_str = sprint_date.strftime("%d.%m.%Y")
    bidding_sheet_name = f"{date_str} - BIDDING SHEET"
    status_sheet_name = f"{date_str} - TASK STATUS"
    summary_sheet_name = f"{date_str} - DAILY SUMMARY"

    # Delete sheets if they already exist
    for name in [bidding_sheet_name, status_sheet_name, summary_sheet_name]:
        if name in workbook.sheetnames:
            del workbook[name]

    # Create sheets
    bidding_sheet = workbook.create_sheet(bidding_sheet_name)
    status_sheet = workbook.create_sheet(status_sheet_name)
    summary_sheet = workbook.create_sheet(summary_sheet_name)

    # --- SHEET 1: BIDDING SHEET ---
    bidding_sheet.append([
        "STUDENT NAME",
        "EMAIL ID",
        "NO OF DAILY TASK PLANNED",
        "LIST OF DAILY TASK ID PLANNED",
        "NO OF WEEKLY TASK PLANNED",
        "LIST OF WEEKLY TASK ID PLANNED",
        "NO OF MONTHLY TASK PLANNED",
        "LIST OF MONTHLY TASK ID PLANNED"
    ])

    students = User.query.filter_by(role="student").order_by(User.name.asc()).all()
    for student in students:
        bid = EliteSprintBid.query.filter_by(session_id=session.id, student_id=student.id).first()
        bidding_sheet.append([
            student.name.upper(),
            student.email,
            bid.daily_count if (bid and bid.daily_count > 0) else "-",
            ", ".join(bid.daily_tasks) if (bid and bid.daily_tasks) else "-",
            bid.weekly_count if (bid and bid.weekly_count > 0) else "-",
            ", ".join(bid.weekly_tasks) if (bid and bid.weekly_tasks) else "-",
            bid.monthly_count if (bid and bid.monthly_count > 0) else "-",
            ", ".join(bid.monthly_tasks) if (bid and bid.monthly_tasks) else "-"
        ])

    # --- SHEET 2: TASK STATUS ---
    daily_tasks = Task.query.filter_by(task_type="daily", status="active").order_by(Task.task_code.asc(), Task.id.asc()).all()
    weekly_tasks = Task.query.filter_by(task_type="weekly", status="active").order_by(Task.task_code.asc(), Task.id.asc()).all()
    monthly_tasks = Task.query.filter_by(task_type="monthly", status="active").order_by(Task.task_code.asc(), Task.id.asc()).all()

    header_row = ["STUDENT NAME", "EMAIL ID"]
    for t in daily_tasks:
        header_row.append(t.task_code or f"TD{t.id}")
    for t in weekly_tasks:
        header_row.append(t.task_code or f"TW{t.id}")
    for t in monthly_tasks:
        header_row.append(t.task_code or f"TM{t.id}")
    header_row += ["BONUS POINTS", "PENALTY POINTS", "TODAY POINTS", "PREVIOUS DAY BALANCE", "MONTHLY TOTAL POINTS"]

    status_sheet.append(header_row)

    day_start = datetime.combine(session.sprint_date, datetime.min.time())
    day_end = datetime.combine(session.sprint_date, datetime.max.time())
    month_start = datetime(sprint_date.year, sprint_date.month, 1)

    for student in students:
        bid = EliteSprintBid.query.filter_by(session_id=session.id, student_id=student.id).first()

        # Check submissions for this student on this day
        submissions = Submission.query.filter_by(student_id=student.id, submission_date=session.sprint_date).all()
        submission_by_task = {}
        for sub in submissions:
            submission_by_task[sub.task_id] = sub.status

        row = [student.name.upper(), student.email]

        # Helper to get bid list
        bid_tasks = []
        if bid:
            bid_tasks = (bid.daily_tasks or []) + (bid.weekly_tasks or []) + (bid.monthly_tasks or [])
            bid_tasks = [t.strip().upper() for t in bid_tasks if t.strip()]

        # Append status for each task type in order
        for tasks_list in [daily_tasks, weekly_tasks, monthly_tasks]:
            for task in tasks_list:
                status_val = submission_by_task.get(task.id)
                is_bidded = (task.task_code and task.task_code.strip().upper() in bid_tasks) or (str(task.id) in bid_tasks)

                if status_val == "approved":
                    symbol = "✅"
                elif is_bidded and status_val != "approved":
                    symbol = "⚠️"
                elif status_val == "rejected":
                    symbol = "❌"
                else:
                    symbol = "-"
                row.append(symbol)

        # Calculate points
        bonus = db.session.query(func.coalesce(func.sum(PointTransaction.points), 0)).filter(
            PointTransaction.student_id == student.id,
            PointTransaction.type == "award",
            PointTransaction.task_id == None,
            PointTransaction.created_at >= day_start,
            PointTransaction.created_at <= day_end
        ).scalar() or 0

        penalty = db.session.query(func.coalesce(func.sum(PointTransaction.points), 0)).filter(
            PointTransaction.student_id == student.id,
            PointTransaction.type == "penalty",
            PointTransaction.created_at >= day_start,
            PointTransaction.created_at <= day_end
        ).scalar() or 0

        today_pts = db.session.query(func.coalesce(func.sum(PointTransaction.points), 0)).filter(
            PointTransaction.student_id == student.id,
            PointTransaction.created_at >= day_start,
            PointTransaction.created_at <= day_end
        ).scalar() or 0

        prev = db.session.query(func.coalesce(func.sum(PointTransaction.points), 0)).filter(
            PointTransaction.student_id == student.id,
            PointTransaction.created_at < day_start
        ).scalar() or 0

        monthly_total = db.session.query(func.coalesce(func.sum(PointTransaction.points), 0)).filter(
            PointTransaction.student_id == student.id,
            PointTransaction.created_at >= month_start,
            PointTransaction.created_at <= day_end
        ).scalar() or 0

        row += [bonus, abs(penalty), today_pts, prev, monthly_total]
        status_sheet.append(row)

    # --- SHEET 3: DAILY SUMMARY ---
    total_students = len(students)
    participated_students = EliteSprintBid.query.filter_by(session_id=session.id).count()
    total_daily_planned = db.session.query(func.coalesce(func.sum(EliteSprintBid.daily_count), 0)).filter_by(session_id=session.id).scalar() or 0
    total_weekly_planned = db.session.query(func.coalesce(func.sum(EliteSprintBid.weekly_count), 0)).filter_by(session_id=session.id).scalar() or 0
    total_monthly_planned = db.session.query(func.coalesce(func.sum(EliteSprintBid.monthly_count), 0)).filter_by(session_id=session.id).scalar() or 0

    total_approved = Submission.query.filter_by(submission_date=session.sprint_date, status="approved").count()
    total_rejected = Submission.query.filter_by(submission_date=session.sprint_date, status="rejected").count()
    total_pending = Submission.query.filter_by(submission_date=session.sprint_date, status="waiting_approval").count()

    total_bonus_points = db.session.query(func.coalesce(func.sum(PointTransaction.points), 0)).filter(
        PointTransaction.type == "award",
        PointTransaction.task_id == None,
        PointTransaction.created_at >= day_start,
        PointTransaction.created_at <= day_end
    ).scalar() or 0

    total_penalty_points = db.session.query(func.coalesce(func.sum(PointTransaction.points), 0)).filter(
        PointTransaction.type == "penalty",
        PointTransaction.created_at >= day_start,
        PointTransaction.created_at <= day_end
    ).scalar() or 0
    total_penalty_points = abs(total_penalty_points)

    golden_stars_awarded = EliteSprintBid.query.filter_by(session_id=session.id, golden_star=True).count()

    penalty_cases = db.session.query(func.count(func.distinct(PointTransaction.student_id))).filter(
        PointTransaction.type == "penalty",
        PointTransaction.created_at >= day_start,
        PointTransaction.created_at <= day_end
    ).scalar() or 0

    # Top Performer of the Day
    top_performer_query = db.session.query(
        User.name,
        func.coalesce(func.sum(PointTransaction.points), 0).label("today_pts")
    ).join(PointTransaction, PointTransaction.student_id == User.id)\
     .filter(PointTransaction.created_at >= day_start, PointTransaction.created_at <= day_end)\
     .group_by(User.id, User.name)\
     .order_by(db.desc("today_pts"))\
     .all()
    if top_performer_query and top_performer_query[0][1] > 0:
        max_pts = top_performer_query[0][1]
        top_performers = [row[0] for row in top_performer_query if row[1] == max_pts]
        top_performer_name = ", ".join(top_performers)
    else:
        top_performer_name = "-"

    # Compliance Rate
    compliance_rate = f"{(golden_stars_awarded / participated_students * 100):.2f}%" if participated_students > 0 else "0.00%"

    summary_sheet.append(["METRIC NAME", "VALUE"])
    metrics_data = [
        ("TOTAL STUDENTS", total_students),
        ("STUDENTS PARTICIPATED IN BIDDING", participated_students),
        ("TOTAL DAILY TASKS PLANNED", total_daily_planned),
        ("TOTAL TASKS APPROVED", total_approved),
        ("TOTAL TASKS REJECTED", total_rejected),
        ("TOTAL BONUS POINTS", total_bonus_points),
        ("TOTAL PENALTY POINTS", total_penalty_points),
        ("GOLDEN STARS AWARDED", golden_stars_awarded),
        ("PENALTY CASES", penalty_cases),
        ("TOP PERFORMER OF THE DAY", top_performer_name),
        ("SPRINT COMPLIANCE RATE", compliance_rate)
    ]
    for key, val in metrics_data:
        summary_sheet.append([key, val])

    # Pie Chart Data (Approved vs Rejected vs Pending)
    summary_sheet.cell(row=1, column=5, value="Status")
    summary_sheet.cell(row=1, column=6, value="Count")
    summary_sheet.cell(row=2, column=5, value="Approved")
    summary_sheet.cell(row=2, column=6, value=total_approved)
    summary_sheet.cell(row=3, column=5, value="Rejected")
    summary_sheet.cell(row=3, column=6, value=total_rejected)
    summary_sheet.cell(row=4, column=5, value="Pending")
    summary_sheet.cell(row=4, column=6, value=total_pending)

    # Bar Chart Data (Top 5 students by Today Points)
    summary_sheet.cell(row=1, column=8, value="Student")
    summary_sheet.cell(row=1, column=9, value="Today Points")
    for idx, (name, pts) in enumerate(top_performer_query[:5], start=2):
        summary_sheet.cell(row=idx, column=8, value=name)
        summary_sheet.cell(row=idx, column=9, value=pts)

    # Column Chart Data (Daily vs Weekly vs Monthly planned task counts)
    summary_sheet.cell(row=1, column=11, value="Sprint Type")
    summary_sheet.cell(row=1, column=12, value="Count")
    summary_sheet.cell(row=2, column=11, value="Daily")
    summary_sheet.cell(row=2, column=12, value=total_daily_planned)
    summary_sheet.cell(row=3, column=11, value="Weekly")
    summary_sheet.cell(row=3, column=12, value=total_weekly_planned)
    summary_sheet.cell(row=4, column=11, value="Monthly")
    summary_sheet.cell(row=4, column=12, value=total_monthly_planned)

    # Add Charts
    pie = PieChart()
    pie.title = "Approved vs Rejected vs Pending"
    labels_pie = Reference(summary_sheet, min_col=5, min_row=2, max_row=4)
    data_pie = Reference(summary_sheet, min_col=6, min_row=1, max_row=4)
    pie.add_data(data_pie, titles_from_data=True)
    pie.set_categories(labels_pie)
    summary_sheet.add_chart(pie, "E7")

    bar = BarChart()
    bar.type = "bar"
    bar.title = "Top 5 Students by Today Points"
    bar.x_axis.title = "Points"
    bar.y_axis.title = "Student"
    if len(top_performer_query) > 0:
        num_rows = min(5, len(top_performer_query))
        labels_bar = Reference(summary_sheet, min_col=8, min_row=2, max_row=num_rows+1)
        data_bar = Reference(summary_sheet, min_col=9, min_row=1, max_row=num_rows+1)
        bar.add_data(data_bar, titles_from_data=True)
        bar.set_categories(labels_bar)
    bar.legend = None
    summary_sheet.add_chart(bar, "E22")

    col = BarChart()
    col.type = "col"
    col.title = "Daily vs Weekly vs Monthly Planned Tasks"
    col.x_axis.title = "Sprint Type"
    col.y_axis.title = "Count"
    labels_col = Reference(summary_sheet, min_col=11, min_row=2, max_row=4)
    data_col = Reference(summary_sheet, min_col=12, min_row=1, max_row=4)
    col.add_data(data_col, titles_from_data=True)
    col.set_categories(labels_col)
    col.legend = None
    summary_sheet.add_chart(col, "E37")

    # --- FORMATTING ---
    DARK_BLUE = "1B365D"
    WHITE = "FFFFFF"

    header_font = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
    header_fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
    cell_font = Font(name="Segoe UI", size=10)

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    green_font = Font(name="Segoe UI", size=10, color="006100")

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(name="Segoe UI", size=10, color="9C0006")

    orange_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    orange_font = Font(name="Segoe UI", size=10, color="9C6500")

    gold_fill = PatternFill(start_color="FFE599", end_color="FFE599", fill_type="solid")
    gold_font = Font(name="Segoe UI", size=10, color="7F6000")

    for sheet in [bidding_sheet, status_sheet, summary_sheet]:
        # 1. Header styling
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 2. Freeze panes at A2
        sheet.freeze_panes = "A2"

        # 3. Auto filter
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

        # 4. Auto adjust column widths & cell font
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
                if cell.row > 1:
                    cell.font = cell_font
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # 5. Conditional formatting for cells (✅, ❌, ⚠️, ⭐)
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            for cell in row:
                val = str(cell.value or "")
                if "✅" in val:
                    cell.fill = green_fill
                    cell.font = green_font
                elif "❌" in val:
                    cell.fill = red_fill
                    cell.font = red_font
                elif "⚠️" in val:
                    cell.fill = orange_fill
                    cell.font = orange_font
                elif "⭐" in val:
                    cell.fill = gold_fill
                    cell.font = gold_font

    workbook.save(filepath)
    return send_file(
        filepath,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@elite_sprint_bp.route("/student", methods=["GET", "POST"])
@role_required("student")
def student():
    active_session = get_active_session()
    latest_session = active_session or get_latest_session()
    bid = None
    if active_session:
        bid = EliteSprintBid.query.filter_by(session_id=active_session.id, student_id=current_user.id).first()

    if request.method == "POST":
        if not active_session:
            flash("Elite Sprint Closed.", "warning")
            return redirect(url_for("elite_sprint.student"))
        if bid and bid.is_locked:
            flash("Your sprint bid is locked and cannot be edited.", "danger")
            return redirect(url_for("elite_sprint.student"))
        try:
            daily_tasks = parse_task_ids("daily")
            weekly_tasks = parse_task_ids("weekly")
            monthly_tasks = parse_task_ids("monthly")
            daily_tasks, weekly_tasks, monthly_tasks = filter_tasks_for_sprint_type(
                active_session.sprint_type,
                daily_tasks,
                weekly_tasks,
                monthly_tasks,
            )
            validate_task_ids(daily_tasks, weekly_tasks, monthly_tasks)

            if bid is None:
                bid = EliteSprintBid(session_id=active_session.id, student_id=current_user.id)
                db.session.add(bid)
            bid.daily_tasks = daily_tasks
            bid.weekly_tasks = weekly_tasks
            bid.monthly_tasks = monthly_tasks
            bid.daily_count = len(daily_tasks)
            bid.weekly_count = len(weekly_tasks)
            bid.monthly_count = len(monthly_tasks)
            bid.is_locked = True
            bid.submitted_at = datetime.utcnow()
            db.session.commit()
            flash("Your sprint bid has been saved and locked.", "success")
            return redirect(url_for("elite_sprint.student"))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "danger")

    return render_template(
        "elite_sprint/student.html",
        active_session=active_session,
        latest_session=latest_session,
        bid=bid,
        visible_sections=visible_sections_for_sprint_type(active_session.sprint_type if active_session else "overall"),
        leaderboard=sprint_leaderboard(latest_session.id if latest_session else None),
    )


@elite_sprint_bp.route("/leaderboard")
@role_required("admin", "student")
def leaderboard():
    session = get_active_session() or get_latest_session()
    return render_template(
        "elite_sprint/leaderboard.html",
        session=session,
        leaderboard=sprint_leaderboard(session.id if session else None),
    )


@elite_sprint_bp.route("/analytics")
@role_required("admin")
def analytics():
    session = get_active_session() or get_latest_session()
    metrics = analytics_metrics(session)
    return render_template("elite_sprint/analytics.html", session=session, metrics=metrics)


@elite_sprint_bp.route("/export")
@role_required("admin")
def export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sprint Bids"
    sheet.append(
        [
            "Session ID",
            "Sprint Date",
            "Sprint Type",
            "Status",
            "Student",
            "Email",
            "Register Number",
            "Department",
            "Daily Count",
            "Daily Tasks",
            "Weekly Count",
            "Weekly Tasks",
            "Monthly Count",
            "Monthly Tasks",
            "Total Sprint Tasks",
            "Submitted At",
        ]
    )
    rows = (
        EliteSprintBid.query.join(EliteSprintSession).join(User)
        .order_by(EliteSprintSession.sprint_date.desc(), EliteSprintBid.submitted_at.asc())
        .all()
    )
    for bid in rows:
        sheet.append(
            [
                bid.session_id,
                bid.session.sprint_date,
                bid.session.sprint_type,
                bid.session.status,
                bid.student.name,
                bid.student.email,
                bid.student.register_number or "",
                bid.student.department or "",
                bid.daily_count,
                ", ".join(map(str, bid.daily_tasks or [])),
                bid.weekly_count,
                ", ".join(map(str, bid.weekly_tasks or [])),
                bid.monthly_count,
                ", ".join(map(str, bid.monthly_tasks or [])),
                bid.total_tasks_bidded,
                bid.submitted_at,
            ]
        )
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBFF")
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 44)
        sheet.column_dimensions[column_cells[0].column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="elite_sprint_bidding.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def parse_task_ids(section):
    expected_count = parse_non_negative_int(request.form.get(f"{section}_count"), section)
    raw_ids = [value.strip() for value in request.form.getlist(f"{section}_tasks") if value.strip()]
    if len(raw_ids) != expected_count:
        raise ValueError(f"{section.title()} section must contain exactly {expected_count} task ID(s).")
    task_ids = [value.upper() for value in raw_ids]
    if len(task_ids) != len(set(task_ids)):
        raise ValueError(f"Duplicate task IDs are not allowed in the {section.title()} section.")
    return task_ids


def parse_non_negative_int(value, section):
    try:
        number = int(value or 0)
    except ValueError as exc:
        raise ValueError(f"{section.title()} task count must be a number.") from exc
    if number < 0:
        raise ValueError(f"{section.title()} task count cannot be negative.")
    if number > 100:
        raise ValueError(f"{section.title()} task count is too high for one sprint.")
    return number


def validate_task_ids(daily_tasks, weekly_tasks, monthly_tasks):
    grouped, _ = active_task_ids_by_type()
    sections = {
        "Daily": (daily_tasks, grouped["daily"]),
        "Weekly": (weekly_tasks, grouped["weekly"]),
        "Monthly": (monthly_tasks, grouped["monthly"]),
    }
    for label, (task_ids, active_ids) in sections.items():
        invalid = sorted(set(task_ids) - active_ids)
        if invalid:
            raise ValueError(f"{label} task IDs must exist as active {label.lower()} tasks: {', '.join(map(str, invalid))}.")


def visible_sections_for_sprint_type(sprint_type):
    if sprint_type in {"daily", "weekly", "monthly"}:
        return {sprint_type}
    return {"daily", "weekly", "monthly"}


def filter_tasks_for_sprint_type(sprint_type, daily_tasks, weekly_tasks, monthly_tasks):
    visible = visible_sections_for_sprint_type(sprint_type)
    return (
        daily_tasks if "daily" in visible else [],
        weekly_tasks if "weekly" in visible else [],
        monthly_tasks if "monthly" in visible else [],
    )


def analytics_metrics(session):
    if not session:
        return {
            "participants": 0,
            "average_tasks": 0,
            "highest_daily": 0,
            "highest_weekly": 0,
            "highest_monthly": 0,
            "trend_labels": [],
            "trend_values": [],
        }

    totals = db.session.query(
        func.count(EliteSprintBid.id),
        func.coalesce(func.avg(EliteSprintBid.daily_count + EliteSprintBid.weekly_count + EliteSprintBid.monthly_count), 0),
        func.coalesce(func.max(EliteSprintBid.daily_count), 0),
        func.coalesce(func.max(EliteSprintBid.weekly_count), 0),
        func.coalesce(func.max(EliteSprintBid.monthly_count), 0),
    ).filter(EliteSprintBid.session_id == session.id).one()
    trend_rows = (
        db.session.query(EliteSprintSession.sprint_date, func.count(EliteSprintBid.id))
        .outerjoin(EliteSprintBid, EliteSprintBid.session_id == EliteSprintSession.id)
        .group_by(EliteSprintSession.sprint_date)
        .order_by(EliteSprintSession.sprint_date.asc())
        .limit(30)
        .all()
    )
    return {
        "participants": int(totals[0] or 0),
        "average_tasks": round(float(totals[1] or 0), 2),
        "highest_daily": int(totals[2] or 0),
        "highest_weekly": int(totals[3] or 0),
        "highest_monthly": int(totals[4] or 0),
        "trend_labels": [row[0].strftime("%d %b") for row in trend_rows],
        "trend_values": [int(row[1] or 0) for row in trend_rows],
    }
